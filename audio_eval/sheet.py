"""Read the evaluation workbook and write the completed one.

Columns are found by header name, so the tool works on any sheet in this
format (any number of rows, columns in any order).
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = {
    "text": ("text",),
    "audio_a": ("audio_a", "audio a"),
    "audio_b": ("audio_b", "audio b"),
    "preference": ("preference",),
    "reason": ("reason",),
    "comment": ("comment",),
}


@dataclass
class Case:
    row: int
    text: str
    audio_a: str
    audio_b: str


class EvalSheet:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.wb = openpyxl.load_workbook(self.path)
        self.ws, self.cols = self._find_worksheet()

    def _find_worksheet(self):
        for ws in self.wb.worksheets:
            header = {str(c.value).strip().lower(): c.column for c in ws[1] if c.value}
            cols = {}
            for key, names in HEADERS.items():
                for h, col in header.items():
                    if any(h.startswith(n) for n in names):
                        cols[key] = col
                        break
            if {"text", "audio_a", "audio_b"} <= cols.keys():
                # answer columns may be missing in a bare sheet: append them
                nxt = ws.max_column + 1
                for key, title in (("preference", "Preference Selection"),
                                   ("reason", "Reason(Multiple-choice)"), ("comment", "Comments")):
                    if key not in cols:
                        ws.cell(1, nxt, title)
                        cols[key] = nxt
                        nxt += 1
                return ws, cols
        raise ValueError("No sheet with 'text', 'Audio_A' and 'Audio_B' header columns found.")

    def cases(self) -> list[Case]:
        out = []
        for r in range(2, self.ws.max_row + 1):
            get = lambda k: self.ws.cell(r, self.cols[k]).value  # noqa: E731
            if get("text") and get("audio_a") and get("audio_b"):
                out.append(Case(r, str(get("text")).strip(), str(get("audio_a")).strip(), str(get("audio_b")).strip()))
        return out

    def write_answer(self, row: int, preference: str, reasons: list[str], comment: str) -> None:
        wrap = Alignment(wrap_text=True, vertical="top")
        for key, value in (("preference", preference), ("reason", "\n".join(reasons)), ("comment", comment)):
            c = self.ws.cell(row, self.cols[key], value)
            c.alignment = wrap
        self.ws.column_dimensions[get_column_letter(self.cols["reason"])].width = 55
        self.ws.column_dimensions[get_column_letter(self.cols["comment"])].width = 80

    def write_table(self, title: str, rows: list[dict]) -> None:
        if title in self.wb.sheetnames:
            del self.wb[title]
        ws = self.wb.create_sheet(title)
        if not rows:
            return
        keys = list(rows[0].keys())
        for j, k in enumerate(keys, 1):
            c = ws.cell(1, j, k)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDE6F0")
        for i, row in enumerate(rows, 2):
            for j, k in enumerate(keys, 1):
                ws.cell(i, j, row.get(k)).alignment = Alignment(wrap_text=True, vertical="top")
        for j, k in enumerate(keys, 1):
            longest = max(len(str(r.get(k) or "")) for r in rows)
            ws.column_dimensions[get_column_letter(j)].width = min(60, max(12, len(k) + 2, longest // 2))
        ws.freeze_panes = "A2"

    def write_summary(self, text: str) -> None:
        title = "Summary"
        if title in self.wb.sheetnames:
            del self.wb[title]
        ws = self.wb.create_sheet(title, index=self.wb.index(self.ws) + 1)
        ws.column_dimensions["A"].width = 120
        for i, para in enumerate(text.split("\n"), 1):
            c = ws.cell(i, 1, para)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if i == 1:
                c.font = Font(bold=True, size=13)
            elif para.endswith(":") or para.startswith("#"):
                c.font = Font(bold=True)

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        self.wb.save(path)
        return path
