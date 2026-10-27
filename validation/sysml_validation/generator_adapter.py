"""Read the wirelist generator's output into a :class:`GeneratorReport`.

Two sources are supported:

* **JSON** — the neutral intermediate schema, where each subsystem maps categories
  to either a list of element names (element-level diff available) or a bare
  integer count (count-only).
* **XLSX** — the generator's actual spreadsheet. The sheet/column layout differs
  per generator, so :func:`from_spreadsheet` takes a small mapping describing where
  the subsystem, category and element-name live. Fill it in once the real
  spreadsheet is in hand.
"""

from __future__ import annotations

import json
from typing import Any, Dict, Iterable, Mapping

from .model import Element, GeneratorReport


def from_json(path: str) -> GeneratorReport:
    """Load the neutral ``{subsystem: {category: [names] | int}}`` schema."""
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    subs = data.get("subsystems", data) if isinstance(data, dict) else {}

    report = GeneratorReport()
    for sub_name, cats in subs.items():
        if not isinstance(cats, Mapping):
            continue
        sub = report.subsystem(sub_name)
        for category, value in cats.items():
            cat = sub.category(category)
            if isinstance(value, int):
                cat.elements = [Element(name=f"<{category}#{i}>") for i in range(value)]
                report.counts_only.add((sub_name, category))
            elif isinstance(value, Iterable):
                cat.elements = [_to_element(v) for v in value]
            else:
                raise ValueError(
                    f"{sub_name}/{category}: expected list or int, got {type(value)!r}"
                )
    return report


def _to_element(value: Any) -> Element:
    """Accept a bare name string or a ``{name, type, id}`` dict."""
    if isinstance(value, str):
        return Element(name=value)
    if isinstance(value, Mapping):
        return Element(
            name=value.get("name") or value.get("@id") or "",
            etype=value.get("type") or value.get("@type") or "",
            eid=value.get("id") or value.get("@id") or "",
        )
    raise ValueError(f"cannot read element from {value!r}")


def from_spreadsheet(
    path: str,
    *,
    sheet: str | int | None = None,
    subsystem_col: str,
    category_col: str,
    name_col: str | None = None,
    category_map: Mapping[str, str] | None = None,
    header_row: int = 1,
) -> GeneratorReport:
    """Read an ``.xlsx`` wirelist into a report.

    Parameters describe the layout: which column holds the subsystem, which the
    category, and (optionally) which the element name. ``category_map`` translates
    the spreadsheet's category labels (e.g. ``"Connection"``) to our canonical
    names (``"connections"``). Rows are grouped; when ``name_col`` is given the
    element-level diff is available, otherwise the cell is count-only.

    Requires ``openpyxl``. The exact column names are generator-specific — wire
    them up once the real spreadsheet is available.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise RuntimeError("from_spreadsheet needs openpyxl installed") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if isinstance(sheet, str) else (
        wb.worksheets[sheet] if isinstance(sheet, int) else wb.active
    )

    rows = ws.iter_rows(values_only=True)
    headers: Dict[str, int] = {}
    for _ in range(header_row):
        header = next(rows)
    headers = {str(h): i for i, h in enumerate(header) if h is not None}

    def col(name: str) -> int:
        if name not in headers:
            raise KeyError(f"column {name!r} not in sheet headers {list(headers)}")
        return headers[name]

    cmap = dict(category_map or {})
    report = GeneratorReport()
    for row in rows:
        if row is None or all(c is None for c in row):
            continue
        sub_name = row[col(subsystem_col)]
        raw_cat = row[col(category_col)]
        if sub_name is None or raw_cat is None:
            continue
        category = cmap.get(str(raw_cat), str(raw_cat))
        sub = report.subsystem(str(sub_name))
        cat = sub.category(category)
        if name_col is not None:
            value = row[col(name_col)]
            cat.elements.append(Element(name="" if value is None else str(value)))
        else:
            # count-only: append an anonymous placeholder per row
            cat.elements.append(Element(name=f"<{category}#{cat.count}>"))
            report.counts_only.add((str(sub_name), category))
    return report
